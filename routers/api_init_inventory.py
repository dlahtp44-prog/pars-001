# app/routers/api_init_inventory.py
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Tuple

import openpyxl
from fastapi import APIRouter, File, Form, HTTPException, UploadFile

from app.db import get_db, upsert_inventory, add_history

router = APIRouter(prefix="/api/init", tags=["초기재고 세팅"])


# =====================================================
# CONFIG
# =====================================================

# ✅ 기존 "입고 엑셀"과 동일 컬럼 (한글 고정)
REQUIRED_COLS = ["창고", "로케이션", "품번", "LOT", "규격", "수량"]
OPTIONAL_COLS = ["브랜드", "품명", "비고"]

ALL_COLS = REQUIRED_COLS + [c for c in OPTIONAL_COLS if c not in REQUIRED_COLS]


def _norm(v: Any) -> str:
    return ("" if v is None else str(v)).strip()


def _q3(v: Any) -> float:
    try:
        d = Decimal(str(v)).quantize(Decimal("0.000"), rounding=ROUND_HALF_UP)
        return float(d)
    except Exception:
        return 0.0


# =====================================================
# EXCEL PARSER
# =====================================================

def _read_excel_rows(data: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Returns:
      ok_rows: [{warehouse, location, brand, item_code, item_name, lot, spec, qty, note}, ...]
      err_rows: [{rownum, error, raw}, ...]
    """
    wb = openpyxl.load_workbook(filename=io.BytesIO(data), data_only=True)
    ws = wb.active

    try:
        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        raise HTTPException(status_code=400, detail="엑셀에 데이터가 없습니다.")

    headers = [_norm(h) for h in header_cells]
    col_index = {h: i for i, h in enumerate(headers) if h}

    missing = [c for c in REQUIRED_COLS if c not in col_index]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"필수 컬럼 누락: {', '.join(missing)} (기존 입고 엑셀 양식과 동일해야 합니다.)",
        )

    ok_rows: List[Dict[str, Any]] = []
    err_rows: List[Dict[str, Any]] = []

    # 데이터 시작: 2행부터
    for ridx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all((_norm(x) == "" for x in row)):
            continue

        def get(col: str) -> str:
            return _norm(row[col_index[col]]) if col in col_index else ""

        raw = {h: ("" if i >= len(row) else _norm(row[i])) for h, i in col_index.items()}

        warehouse = get("창고")
        location = get("로케이션")
        brand = get("브랜드")
        item_code = get("품번")
        item_name = get("품명")
        lot = get("LOT")
        spec = get("규격")
        note = get("비고")

        qty_raw = row[col_index["수량"]]
        qty = _q3(qty_raw)

        if not warehouse or not location or not item_code or not lot or not spec:
            err_rows.append({
                "rownum": ridx,
                "error": "필수값(창고/로케이션/품번/LOT/규격) 누락",
                "raw": raw
            })
            continue

        if qty <= 0:
            err_rows.append({
                "rownum": ridx,
                "error": "수량은 0보다 커야 합니다.",
                "raw": raw
            })
            continue

        ok_rows.append({
            "warehouse": warehouse,
            "location": location,
            "brand": brand,
            "item_code": item_code,
            "item_name": item_name,
            "lot": lot,
            "spec": spec,
            "qty": qty,
            "note": note,
        })

    # 중복 키 체크
    seen = {}
    dedup_ok: List[Dict[str, Any]] = []
    for r in ok_rows:
        key = (r["warehouse"], r["location"], r["brand"], r["item_code"], r["lot"], r["spec"])
        if key in seen:
            err_rows.append({
                "rownum": None,
                "error": f"중복 키 발견: {key}",
                "raw": {"first": seen[key], "dup": r},
            })
            continue
        seen[key] = r
        dedup_ok.append(r)

    return dedup_ok, err_rows


# =====================================================
# DB COUNTS / UTILS
# =====================================================

def _count_inventory() -> int:
    conn = get_db()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM inventory")
        return int(cur.fetchone()[0])
    finally:
        conn.close()


def _count_history() -> int:
    conn = get_db()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM history")
        return int(cur.fetchone()[0])
    finally:
        conn.close()


def _make_batch_id() -> str:
    return "INIT-" + datetime.now().strftime("%Y%m%d-%H%M%S")


# =====================================================
# APIs
# =====================================================

@router.get("/status")
def init_inventory_status():
    """
    🔍 초기재고 상태 조회
    - UI 경고용
    """
    inv_cnt = _count_inventory()
    hist_cnt = _count_history()

    return {
        "inventory_count": inv_cnt,
        "history_count": hist_cnt,
        "has_data": inv_cnt > 0 or hist_cnt > 0,
    }


@router.post("/preview")
async def init_preview(file: UploadFile = File(...)):
    """
    초기재고 엑셀 업로드 미리보기(검증)
    """
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="엑셀(.xlsx) 파일만 업로드 가능합니다.")

    data = await file.read()
    ok_rows, err_rows = _read_excel_rows(data)

    return {
        "ok": True,
        "inventory_count": _count_inventory(),
        "history_count": _count_history(),
        "summary": {
            "total_rows": len(ok_rows) + len(err_rows),
            "ok_rows": len(ok_rows),
            "error_rows": len(err_rows),
        },
        "rows_ok": ok_rows[:2000],
        "rows_error": err_rows[:2000],
        "message": f"총 {len(ok_rows) + len(err_rows)}행 중 정상 {len(ok_rows)}행, 오류 {len(err_rows)}행",
    }


@router.post("/commit")
async def init_commit(
    file: UploadFile = File(...),
    operator: str = Form(...),
    confirm: str = Form(""),
    force: int = Form(0),
):
    """
    초기재고 반영(커밋)
    """
    if confirm.strip() != "INIT-CONFIRM":
        raise HTTPException(status_code=400, detail="확인 문구가 필요합니다. confirm=INIT-CONFIRM")

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="엑셀(.xlsx) 파일만 업로드 가능합니다.")

    inv_cnt = _count_inventory()
    if inv_cnt > 0 and int(force) != 1:
        raise HTTPException(
            status_code=400,
            detail=f"현재 inventory가 {inv_cnt}건 존재합니다. 전체 리셋 후 진행을 권장합니다. (강제: force=1)",
        )

    data = await file.read()
    ok_rows, err_rows = _read_excel_rows(data)

    if not ok_rows:
        raise HTTPException(status_code=400, detail="반영할 정상 데이터가 없습니다.")

    batch_id = _make_batch_id()
    applied = 0
    failed: List[Dict[str, Any]] = []

    for r in ok_rows:
        try:
            ok = upsert_inventory(
                r["warehouse"],
                r["location"],
                r["brand"],
                r["item_code"],
                r["item_name"],
                r["lot"],
                r["spec"],
                r["qty"],
                note=(r.get("note") or "초기재고"),
            )
            if not ok:
                raise ValueError("재고 반영 실패")

            add_history(
                "초기재고",
                r["warehouse"],
                operator,
                r["brand"],
                r["item_code"],
                r["item_name"],
                r["lot"],
                r["spec"],
                "INIT",
                r["location"],
                r["qty"],
                note=(r.get("note") or "초기재고 세팅"),
                batch_id=batch_id,
                dedup_seconds=0,
            )

            applied += 1

        except Exception as e:
            failed.append({"row": r, "error": str(e)})

    return {
        "ok": True,
        "batch_id": batch_id,
        "summary": {
            "total_ok_rows": len(ok_rows),
            "applied": applied,
            "failed": len(failed),
            "error_rows_in_file": len(err_rows),
        },
        "failed_rows": failed[:200],
        "message": f"초기재고 반영 완료: 정상 {len(ok_rows)}행 중 {applied}행 적용",
    }
